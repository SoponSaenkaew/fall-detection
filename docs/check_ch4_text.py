import docx

doc = docx.Document(r"D:\Final project\fall-detection\docs\เอกสารประกอบโปรเจกต์จบ\รายงานเชิงเทคนิก\รายงานเชิงเทคนิค.docx")

print("=== TEXT AROUND TABLE 4.1 AND 4.2 ===")
for i, p in enumerate(doc.paragraphs):
    if "4.5" in p.text or "ตารางที่ 4" in p.text or "Confusion" in p.text or "LD2410" in p.text:
        print(f"\n[P#{i}]")
        print(p.text)
        # print next 3 paragraphs
        for j in range(1, 4):
            if i + j < len(doc.paragraphs):
                print(f"  + {doc.paragraphs[i+j].text}")
