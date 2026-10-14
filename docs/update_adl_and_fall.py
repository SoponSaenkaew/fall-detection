import os
import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def set_run_font(run, font_name="TH SarabunPSK", font_size=16, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)

def set_cell_text(cell, text, font_name="TH SarabunPSK", font_size=16, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    set_run_font(run, font_name=font_name, font_size=font_size, bold=bold)

def set_paragraph_text(p, text, font_name="TH SarabunPSK", font_size=16, bold=False, align=None):
    p.text = ""
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, font_name=font_name, font_size=font_size, bold=bold)

def format_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def update_chapter3(path):
    print(f"\nUpdating Chapter 3 in {path}...")
    doc = docx.Document(path)
    
    # 3.7 Intro paragraph
    for p in doc.paragraphs:
        p_clean = p.text.replace('\u200b', '').replace('\ufeff', '')
        if '3.7' in p_clean and 'การออกแบบการทดสอบ' in p_clean:
            # find next paragraphs
            pass
        if 'กำหนดแผนการทดสอบครอบคลุมทั้งท่าทางการหกล้มและกิจกรรมในชีวิตประจำวัน' in p_clean:
            new_text = "เพื่อให้การประเมินประสิทธิภาพของระบบมีความน่าเชื่อถือและสะท้อนสภาพการใช้งานจริงในห้องน้ำ การวิจัยนี้จึงได้ออกแบบการทดลองเชิงประจักษ์ โดยทำการติดตั้งชุดอุปกรณ์ต้นแบบที่ระดับความสูง 2.2 เมตร บริเวณกึ่งกลางเพดานของห้องน้ำทดสอบขนาด 2.0 × 2.0 เมตร และกำหนดแผนการทดสอบครอบคลุมทั้งท่าทางการหกล้มและกิจกรรมในชีวิตประจำวัน รวมทั้งสิ้น 120 ครั้งทดสอบ"
            set_paragraph_text(p, new_text)
            print("  Updated 3.7 intro text.")
        
        if '3.7.1' in p_clean and 'การกำหนดรูปแบบการทดสอบ' in p_clean:
            pass
        if 'การทดสอบแบ่งออกเป็น 4 รูปแบบหลัก' in p_clean or 'จำลองการลื่นไถลหรือสะดุดขา' in p_clean:
            new_text = "การทดสอบแบ่งออกเป็น 2 ส่วนหลัก รวมทั้งสิ้น 120 ครั้งทดสอบ ได้แก่ (1) การจำลองเหตุการณ์การหกล้ม 4 ท่าทางหลัก (จำนวนรวม 80 ครั้ง ท่าทางละ 20 ครั้ง) ได้แก่ การล้มไปข้างหน้า (Forward Fall: 20 ครั้ง), การล้มหงายหลัง (Backward Fall: 20 ครั้ง), การล้มไปด้านข้าง (Lateral Fall: 20 ครั้ง) และการล้มทรุดตัวลงกับพื้น (Slow Collapse: 20 ครั้ง) และ (2) การจำลองกิจกรรมปกติในชีวิตประจำวัน (Activities of Daily Living: ADL จำนวนรวม 40 ครั้ง กิจกรรมละ 10 ครั้ง) ได้แก่ การนั่งชักโครกหรือเก้าอี้ (Sitting: 10 ครั้ง), การเดินไปมาหรือเดินเข้าห้องน้ำ (Walking: 10 ครั้ง), การนั่งยองๆ (Crouching: 10 ครั้ง) และการก้มเก็บของบนพื้น (Bending Down: 10 ครั้ง) เพื่อประเมินความสามารถของระบบในการป้องกันการเกิดสัญญาณเตือนที่ผิดพลาด (False Alarm Rejection)"
            set_paragraph_text(p, new_text)
            print("  Updated 3.7.1 protocol text.")
            
        if '3.7.2' in p_clean and 'ตัวแปรที่ใช้ในการวิจัย' in p_clean:
            pass
        if 'ตัวแปรต้น (Independent Variables) คือ รูปแบบของท่าทาง' in p_clean:
            new_text = "ตัวแปรที่ใช้ในการทดลองวิจัยแบ่งออกเป็น 3 ด้านหลัก ได้แก่ (1) ตัวแปรต้น (Independent Variables) คือ รูปแบบของท่าทางและพฤติกรรมของมนุษย์ ได้แก่ ท่าทางการล้ม 4 รูปแบบ (ล้มไปข้างหน้า, ล้มหงายหลัง, ล้มด้านข้าง, ล้มทรุดตัว) และกิจกรรมปกติในชีวิตประจำวัน (ADL: นั่งชักโครก, เดินไปมา, นั่งยองๆ, ก้มเก็บของ) (2) ตัวแปรตาม (Dependent Variables) คือ ความถูกต้องในการจำแนกเหตุการณ์ของระบบ (Sensitivity, Specificity, Accuracy), ความสามารถในการคัดกรองสัญญาณเตือนลวง และเวลาหน่วงในการตอบสนองแจ้งเตือน (Radar Confirm Time) และ (3) ตัวแปรควบคุม (Controlled Variables) คือ ความสูงในการติดตั้งเซนเซอร์เรดาร์ (2.2 เมตร), ขนาดพื้นที่ห้องน้ำทดสอบ (2.0 × 2.0 เมตร), แหล่งจ่ายไฟระบบ และค่าเกณฑ์พารามิเตอร์การตรวจจับของเรดาร์"
            set_paragraph_text(p, new_text)
            print("  Updated 3.7.2 variables text.")
            
    doc.save(path)
    print(f"Saved Chapter 3 updates to {path}")

def update_chapter4_and_full(path):
    print(f"\nUpdating Chapter 4 in {path}...")
    doc = docx.Document(path)
    
    # 1. Update text paragraphs
    for i, p in enumerate(doc.paragraphs):
        p_clean = p.text.replace('\u200b', '').replace('\ufeff', '')
        
        # P45 intro to results
        if 'การทดสอบและประเมินประสิทธิภาพของระบบดำเนินการในห้องน้ำทดสอบจำลอง' in p_clean and 'Confusion Matrix' in p_clean:
            new_p45 = "การทดสอบและประเมินประสิทธิภาพของระบบดำเนินการในห้องน้ำทดสอบจำลองขนาด 2.0 × 2.0 เมตร ที่ระดับความสูงติดตั้งเซนเซอร์ 2.2 เมตร โดยบันทึกข้อมูลการทดลองจริงลงในไฟล์สถิติ CSV ซึ่งประกอบด้วยการจำลองเหตุการณ์ล้ม 4 ท่าทางหลัก จำนวนรวม 80 ครั้งทดสอบ (ล้มไปข้างหน้า 20 ครั้ง, ล้มหงายหลัง 20 ครั้ง, ล้มด้านข้าง 20 ครั้ง และล้มทรุดตัวลงกับพื้น 20 ครั้ง) และการทดสอบกิจกรรมปกติในชีวิตประจำวัน (Activities of Daily Living: ADL จำนวน 40 ครั้ง ประกอบด้วยการนั่งชักโครก/เก้าอี้ 10 ครั้ง, การเดินไปมา/เข้าห้องน้ำ 10 ครั้ง, การนั่งยองๆ 10 ครั้ง และการก้มเก็บของ 10 ครั้ง) รวมทั้งสิ้น 120 ครั้งทดสอบ เพื่อนำผลการทดสอบเชิงสถิติมาบันทึกลงในเมทริกซ์ความสับสน (Confusion Matrix) และประเมินค่าตัวชี้วัดประสิทธิภาพเชิงปริมาณตามที่กำหนดไว้ในตารางที่ 4.1"
            set_paragraph_text(p, new_p45)
            print(f"  Updated results intro paragraph at P{i}")
            
        # P48 discussion of results
        if 'จากผลการทดสอบเชิงประจักษ์ที่บันทึกผ่านระบบเก็บสถิติ CSV ในตารางที่ 4.1' in p_clean:
            new_p48 = "จากผลการทดสอบเชิงประจักษ์ที่บันทึกผ่านระบบเก็บสถิติ CSV ในตารางที่ 4.1 พบว่าระบบมีค่าความไวในการตรวจจับการล้ม (Sensitivity / Recall) สูงถึง 100.0% โดยสามารถตรวจจับเหตุการณ์ล้มได้ถูกต้องครบถ้วนทั้ง 80 ครั้ง จากการทดสอบทั้งหมด 4 ท่าทางหลัก (ล้มไปข้างหน้า 20 ครั้ง, ล้มหงายหลัง 20 ครั้ง, ล้มด้านข้าง 20 ครั้ง และล้มทรุดตัว 20 ครั้ง) โดยไม่มีการตรวจจับผิดพลาดหรือหลุดรอดเลยแม้แต่ครั้งเดียว (True Positive = 80 ครั้ง, False Negative = 0 ครั้ง) และมีค่าเวลาหน่วงในการประมวลผลยืนยันการล้มของเรดาร์ (Confirm Time) รวดเร็วเฉลี่ยเพียง 37.45 มิลลิวินาที (อยู่ในช่วง 34 ถึง 52 มิลลิวินาที) ในขณะที่การทดสอบกิจกรรมปกติในชีวิตประจำวัน (ADL) ทั้ง 4 รูปแบบ จำนวนรวม 40 ครั้ง ระบบสามารถจำแนกสภาวะการเคลื่อนไหวปกติได้อย่างแม่นยำโดยไม่เกิดสัญญาณแจ้งเตือนผิดพลาดเลยแม้แต่ครั้งเดียว (True Negative = 40 ครั้ง, False Positive = 0 ครั้ง) คิดเป็นค่าความจำเพาะต่อกิจกรรมปกติ (Specificity) สูงถึง 100.0% ส่งผลให้ระบบมีความแม่นยำในการทำนาย (Precision / PPV) อยู่ที่ 100.0% และมีความถูกต้องโดยรวมของระบบ (Overall Accuracy) สมบูรณ์แบบที่ 100.0% (120 จาก 120 ครั้งทดสอบ) ดังรายละเอียดผลการจำแนกกิจกรรมปกติแต่ละรูปแบบที่แสดงในตารางที่ 4.2"
            set_paragraph_text(p, new_p48)
            print(f"  Updated results discussion paragraph at P{i}")

        # Update references to Table 4.2 -> Table 4.3 for 24 GHz presence
        if 'การประเมินประสิทธิภาพของโมดูลเรดาร์ 24 GHz' in p_clean and 'ตารางที่ 4.2' in p_clean:
            new_t = p_clean.replace('ตารางที่ 4.2', 'ตารางที่ 4.3')
            set_paragraph_text(p, new_t)
            print(f"  Updated 24GHz presence intro reference at P{i}")
            
        if 'ตารางที่ 4.2  ผลการทดสอบประสิทธิภาพการตรวจจับการมีอยู่ของบุคคลด้วยเรดาร์ 24 GHz' in p_clean:
            new_t = p_clean.replace('ตารางที่ 4.2', 'ตารางที่ 4.3')
            set_paragraph_text(p, new_t, bold=True)
            print(f"  Updated 24GHz table title at P{i}")
            
        if 'จากผลการทดสอบในตารางที่ 4.2 พบว่าโมดูลเรดาร์ 24 GHz' in p_clean:
            new_t = p_clean.replace('ตารางที่ 4.2', 'ตารางที่ 4.3')
            set_paragraph_text(p, new_t)
            print(f"  Updated 24GHz discussion reference at P{i}")

    # 2. Update Table 4.1
    for t in doc.tables:
        first_row_txt = [c.text.strip().replace('\u200b', '') for c in t.rows[0].cells]
        if any('ประเภทการทดสอบ' in c for c in first_row_txt) and len(t.columns) == 3:
            print("  Found Table 4.1! Updating rows...")
            # We need 17 rows (header + 5 fall + 5 adl + 6 metrics)
            # If currently 16 rows, insert 1 row before row 10
            if len(t.rows) == 16:
                t.add_row()
                # Shift rows from index 10 downward
                for r_idx in range(len(t.rows)-1, 10, -1):
                    for c_idx in range(3):
                        t.rows[r_idx].cells[c_idx].text = t.rows[r_idx-1].cells[c_idx].text
            
            # Row 1: Fall header
            set_cell_text(t.rows[1].cells[0], 'เหตุการณ์จำลองการล้มจริง (Fall Event: 80 ครั้ง)', bold=True)
            set_cell_text(t.rows[1].cells[1], '80 ครั้ง (True Positive: TP)', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
            set_cell_text(t.rows[1].cells[2], '0 ครั้ง (False Negative: FN)', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
            
            # Row 2: Forward
            set_cell_text(t.rows[2].cells[0], '- ท่าทางล้มไปข้างหน้า (Forward Fall: 20 ครั้ง)')
            set_cell_text(t.rows[2].cells[1], '20 ครั้ง (100.0%)', align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(t.rows[2].cells[2], '0 ครั้ง (0.0%)', align=WD_ALIGN_PARAGRAPH.CENTER)
            
            # Row 3: Backward (20 trials)
            set_cell_text(t.rows[3].cells[0], '- ท่าทางล้มหงายหลัง (Backward Fall: 20 ครั้ง)')
            set_cell_text(t.rows[3].cells[1], '20 ครั้ง (100.0%)', align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(t.rows[3].cells[2], '0 ครั้ง (0.0%)', align=WD_ALIGN_PARAGRAPH.CENTER)
            
            # Row 4: Lateral
            set_cell_text(t.rows[4].cells[0], '- ท่าทางล้มไปด้านข้าง (Lateral Fall: 20 ครั้ง)')
            set_cell_text(t.rows[4].cells[1], '20 ครั้ง (100.0%)', align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(t.rows[4].cells[2], '0 ครั้ง (0.0%)', align=WD_ALIGN_PARAGRAPH.CENTER)
            
            # Row 5: Slow collapse
            set_cell_text(t.rows[5].cells[0], '- ท่าทางล้มทรุดตัวลงกับพื้น (Slow Collapse: 20 ครั้ง)')
            set_cell_text(t.rows[5].cells[1], '20 ครั้ง (100.0%)', align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(t.rows[5].cells[2], '0 ครั้ง (0.0%)', align=WD_ALIGN_PARAGRAPH.CENTER)
            
            # Row 6: ADL Header
            set_cell_text(t.rows[6].cells[0], 'กิจกรรมปกติในชีวิตประจำวัน (ADL Event: 40 ครั้ง)', bold=True)
            set_cell_text(t.rows[6].cells[1], '0 ครั้ง (False Positive: FP)', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
            set_cell_text(t.rows[6].cells[2], '40 ครั้ง (True Negative: TN)', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
            
            # Row 7: Sitting
            set_cell_text(t.rows[7].cells[0], '- นั่งชักโครก / เก้าอี้ (Sitting: 10 ครั้ง)')
            set_cell_text(t.rows[7].cells[1], '0 ครั้ง (0.0%)', align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(t.rows[7].cells[2], '10 ครั้ง (100.0%)', align=WD_ALIGN_PARAGRAPH.CENTER)
            
            # Row 8: Walking
            set_cell_text(t.rows[8].cells[0], '- เดินไปมา / เข้าห้องน้ำ (Walking: 10 ครั้ง)')
            set_cell_text(t.rows[8].cells[1], '0 ครั้ง (0.0%)', align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(t.rows[8].cells[2], '10 ครั้ง (100.0%)', align=WD_ALIGN_PARAGRAPH.CENTER)
            
            # Row 9: Crouching
            set_cell_text(t.rows[9].cells[0], '- นั่งยองๆ (Crouching: 10 ครั้ง)')
            set_cell_text(t.rows[9].cells[1], '0 ครั้ง (0.0%)', align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(t.rows[9].cells[2], '10 ครั้ง (100.0%)', align=WD_ALIGN_PARAGRAPH.CENTER)
            
            # Row 10: Bending down
            set_cell_text(t.rows[10].cells[0], '- ก้มเก็บของ (Bending Down: 10 ครั้ง)')
            set_cell_text(t.rows[10].cells[1], '0 ครั้ง (0.0%)', align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(t.rows[10].cells[2], '10 ครั้ง (100.0%)', align=WD_ALIGN_PARAGRAPH.CENTER)
            
            # Row 11: Metrics header
            set_cell_text(t.rows[11].cells[0], 'ตัวชี้วัดประสิทธิภาพเชิงปริมาณ (Performance Metrics)', bold=True)
            set_cell_text(t.rows[11].cells[1], 'สูตรการคำนวณ', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
            set_cell_text(t.rows[11].cells[2], 'ค่าที่ได้จากการประเมิน', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
            
            # Row 12: Sensitivity
            set_cell_text(t.rows[12].cells[0], 'ความไวในการตรวจจับ (Sensitivity / Recall)')
            set_cell_text(t.rows[12].cells[1], 'TP / (TP + FN)', align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(t.rows[12].cells[2], '100.0% (80 / 80)', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
            
            # Row 13: Specificity
            set_cell_text(t.rows[13].cells[0], 'ความจำเพาะต่อกิจกรรมปกติ (Specificity)')
            set_cell_text(t.rows[13].cells[1], 'TN / (TN + FP)', align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(t.rows[13].cells[2], '100.0% (40 / 40)', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
            
            # Row 14: Precision
            set_cell_text(t.rows[14].cells[0], 'ความแม่นยำในการทำนาย (Precision / PPV)')
            set_cell_text(t.rows[14].cells[1], 'TP / (TP + FP)', align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(t.rows[14].cells[2], '100.0% (80 / 80)', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
            
            # Row 15: Accuracy
            set_cell_text(t.rows[15].cells[0], 'ความถูกต้องโดยรวมของระบบ (Overall Accuracy)')
            set_cell_text(t.rows[15].cells[1], '(TP + TN) / Total', align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(t.rows[15].cells[2], '100.0% (120 / 120)', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
            
            # Row 16: Confirm time
            set_cell_text(t.rows[16].cells[0], 'เวลาหน่วงการยืนยันผลการล้ม (Confirm Time)')
            set_cell_text(t.rows[16].cells[1], 'เวลาประมวลผลเฉลี่ยจากเรดาร์', align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(t.rows[16].cells[2], '37.45 ms (เฉลี่ยจากข้อมูล CSV 80 ครั้ง)', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
            
            break
            
    # 3. Insert dedicated ADL table (Table 4.2)
    # Check if Table 4.2 already inserted
    already_inserted = False
    for t in doc.tables:
        for r in t.rows:
            txt = ' '.join(c.text for c in r.cells)
            if 'นั่งชักโครก / เก้าอี้ (Sitting)' in txt and len(r.cells) == 6:
                already_inserted = True
                break
                
    if not already_inserted:
        print("  Inserting dedicated 6-column ADL Table 4.2...")
        # Find the discussion paragraph (P48)
        target_p = None
        for p in doc.paragraphs:
            p_clean = p.text.replace('\u200b', '').replace('\ufeff', '')
            if 'จากผลการทดสอบเชิงประจักษ์ที่บันทึกผ่านระบบเก็บสถิติ CSV ในตารางที่ 4.1' in p_clean:
                target_p = p
                break
                
        if target_p:
            # Create caption paragraph
            p_caption = doc.add_paragraph()
            set_paragraph_text(p_caption, "ตารางที่ 4.2  ผลการทดสอบการจำแนกกิจกรรมปกติในชีวิตประจำวัน (Activities of Daily Living: ADL)", bold=True)
            
            # Create new 6-column table
            adl_table = doc.add_table(rows=6, cols=6)
            adl_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            format_table_borders(adl_table)
            
            headers = ['รูปแบบกิจกรรมปกติ (ADL)', 'จำนวน (ครั้ง)', 'สถานะที่ควรเป็น', 'ผลตรวจจับ: ปกติ', 'ผลตรวจจับ: เตือนลวง', 'อัตราความถูกต้อง']
            for c_idx, h in enumerate(headers):
                set_cell_text(adl_table.rows[0].cells[c_idx], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
                
            adl_data = [
                ['นั่งชักโครก / เก้าอี้ (Sitting)', '10', 'ปกติ', '10', '0', '100%'],
                ['เดินไปมา / เข้าห้องน้ำ (Walking)', '10', 'ปกติ', '10', '0', '100%'],
                ['นั่งยองๆ (Crouching)', '10', 'ปกติ', '10', '0', '100%'],
                ['ก้มเก็บของ (Bending Down)', '10', 'ปกติ', '10', '0', '100%'],
                ['รวมกิจกรรมปกติทั้งหมด (Overall ADL)', '40', 'ปกติ', '40 (100%)', '0 (0%)', '100%']
            ]
            
            for r_idx, row in enumerate(adl_data, start=1):
                is_bold = (r_idx == 5)
                set_cell_text(adl_table.rows[r_idx].cells[0], row[0], bold=is_bold, align=WD_ALIGN_PARAGRAPH.LEFT)
                for c_idx in range(1, 6):
                    set_cell_text(adl_table.rows[r_idx].cells[c_idx], row[c_idx], bold=is_bold, align=WD_ALIGN_PARAGRAPH.CENTER)
                    
            # Insert into document position after target_p
            target_p._p.addnext(adl_table._tbl)
            target_p._p.addnext(p_caption._p)
            print("  Successfully positioned Table 4.2 after discussion paragraph!")
            
    doc.save(path)
    print(f"Saved Chapter 4 updates to {path}")

if __name__ == '__main__':
    update_chapter3('docs/เอกสารประกอบโปรเจกต์จบ/รายงานเชิงเทคนิก/06 บทที่ 3.docx')
    update_chapter4_and_full('docs/เอกสารประกอบโปรเจกต์จบ/รายงานเชิงเทคนิก/07 บทที่ 4.docx')
    update_chapter3('docs/เอกสารประกอบโปรเจกต์จบ/รายงานเชิงเทคนิก/รายงานเชิงเทคนิค_ฉบับสมบูรณ์.docx')
    update_chapter4_and_full('docs/เอกสารประกอบโปรเจกต์จบ/รายงานเชิงเทคนิก/รายงานเชิงเทคนิค_ฉบับสมบูรณ์.docx')
    print("\nAll updates completed successfully!")
