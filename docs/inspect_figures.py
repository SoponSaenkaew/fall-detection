import docx
import re
import os

doc_path = r"D:\Final project\fall-detection\docs\เอกสารประกอบโปรเจกต์จบ\รายงานเชิงเทคนิก\รายงานเชิงเทคนิค.docx"
doc = docx.Document(doc_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")
for i, p in enumerate(doc.paragraphs):
    xml = p._p.xml
    txt = p.text.strip()
    if "<w:drawing>" in xml:
        r_embed = re.findall(r'r:embed="([^"]+)"', xml)
        targets = [doc.part.rels[r].target_ref for r in r_embed if r in doc.part.rels]
        print(f"P[{i:03d}] [DRAWING targets={targets}]: {txt[:60]}")
    elif "รูปที่" in txt or "ตารางที่" in txt:
        print(f"P[{i:03d}] [CAPTION]: {txt[:90]}")
