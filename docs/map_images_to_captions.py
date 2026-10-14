import docx

doc = docx.Document(r"D:\Final project\fall-detection\docs\เอกสารประกอบโปรเจกต์จบ\รายงานเชิงเทคนิก\รายงานเชิงเทคนิค.docx")

print("=== PARAGRAPHS CONTAINING IMAGES OR CAPTIONS ===")
for p_idx, p in enumerate(doc.paragraphs):
    # check if p has drawing
    has_image = "w:drawing" in p._p.xml
    txt = p.text.strip()
    if has_image:
        print(f"P#{p_idx}: [HAS DRAWING] text='{txt}'")
    elif "รูปที่" in txt:
        print(f"P#{p_idx}: [CAPTION] text='{txt}'")

for t_idx, t in enumerate(doc.tables):
    for r_idx, row in enumerate(t.rows):
        for c_idx, cell in enumerate(row.cells):
            for cp_idx, cp in enumerate(cell.paragraphs):
                if "w:drawing" in cp._p.xml:
                    print(f"Table {t_idx} [R{r_idx}C{c_idx}] P#{cp_idx}: [HAS DRAWING] text='{cp.text.strip()}'")
                elif "รูปที่" in cp.text or "พื้นที่สำหรับวางภาพ" in cp.text:
                    print(f"Table {t_idx} [R{r_idx}C{c_idx}] P#{cp_idx}: [PLACEHOLDER/CAPTION] text='{cp.text.strip()}'")
