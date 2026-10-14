import fitz  # PyMuPDF
import re
import os
import docx

pdf_path = r"D:\Final project\fall-detection\docs\เอกสารประกอบโปรเจกต์จบ\รายงานเชิงเทคนิก\รายงานเชิงเทคนิค_final_render.pdf"
docx_path = r"D:\Final project\fall-detection\docs\เอกสารประกอบโปรเจกต์จบ\รายงานเชิงเทคนิก\รายงานเชิงเทคนิค.docx"

print("=" * 60)
print("1. INSPECTING PDF PAGES")
print("=" * 60)
doc = fitz.open(pdf_path)
print(f"Total pages: {len(doc)}")

pages_info = []

for i, page in enumerate(doc):
    text = page.get_text()
    lines = [line.strip() for line in text.split('\n') if line.strip()]
    first_few = ' | '.join(lines[:3]) if lines else '[BLANK PAGE]'
    
    # Check bounding boxes in top area (y < 100 pt) to see page numbers
    blocks = page.get_text("blocks")
    top_blocks = [b for b in blocks if b[1] < 100]
    top_texts = [b[4].strip() for b in top_blocks]
    
    pages_info.append({
        "page_num": i + 1,
        "first_line": lines[0] if lines else "[EMPTY]",
        "snippet": first_few[:70],
        "top_blocks": top_texts
    })

for p in pages_info:
    tb = " // ".join(p['top_blocks'])
    print(f"P.{p['page_num']:02d}: [Top: {tb[:30]}] -> {p['snippet']}")

print("\n" + "=" * 60)
print("2. SEARCHING FOR SUSPICIOUS PLACEHOLDERS")
print("=" * 60)
placeholders = ["TODO", "FIXME", "XXX", "รอใส่", "ใส่ข้อมูล", "Error!", "???", "ไม่พบแหล่งอ้างอิง"]
found_placeholders = []
for i, page in enumerate(doc):
    text = page.get_text()
    for ph in placeholders:
        if ph.lower() in text.lower():
            found_placeholders.append((i + 1, ph))

if found_placeholders:
    print("Found potential placeholders:")
    for p, ph in found_placeholders:
        print(f"  Page {p}: '{ph}'")
else:
    print("No suspicious placeholders found (TODO, FIXME, XXX, etc.)!")

print("\n" + "=" * 60)
print("3. CHECKING TABLE OF CONTENTS VS ACTUAL PAGES")
print("=" * 60)
# Let's inspect docx headings & page breaks
doc_docx = docx.Document(docx_path)
print(f"Total paragraphs in docx: {len(doc_docx.paragraphs)}")
print(f"Total tables in docx: {len(doc_docx.tables)}")
print(f"Total sections in docx: {len(doc_docx.sections)}")

for idx, sec in enumerate(doc_docx.sections):
    print(f"Section {idx}: Top={sec.top_margin.cm:.2f}cm, Bottom={sec.bottom_margin.cm:.2f}cm, Left={sec.left_margin.cm:.2f}cm, Right={sec.right_margin.cm:.2f}cm, Header={sec.header_distance.cm:.2f}cm, DiffFirst={sec.different_first_page_header_footer}")
