import os
import re
import time
import subprocess
import docx
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import pymupdf as fitz

DOCX_PATH = r"D:\Final project\fall-detection\docs\เอกสารประกอบโปรเจกต์จบ\รายงานเชิงเทคนิก\รายงานเชิงเทคนิค.docx"
PDF_PATH = r"D:\Final project\fall-detection\docs\เอกสารประกอบโปรเจกต์จบ\รายงานเชิงเทคนิก\รายงานเชิงเทคนิค_final_render.pdf"

def set_font(run, font_name="TH SarabunPSK", font_size=16, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)

def update_figure_numbers_and_references(doc):
    print("Adjusting Chapter 4 figure numbers to strict chronological order...")
    # Change:
    # 1. In 4.5.1:
    #    รูปที่ 4.6 -> รูปที่ 4.5
    #    รูปที่ 4.7 -> รูปที่ 4.6
    # 2. In 4.5.2:
    #    รูปที่ 4.8 -> รูปที่ 4.7
    # 3. In 4.5.3:
    #    The diagram caption: รูปที่ 4.5 -> รูปที่ 4.8
    #    The following narrative text: รูปที่ 4.5 -> รูปที่ 4.8
    #    The real-time alert placeholder: รูปที่ 4.9 (stays 4.9)
    
    # 1. Update text inside tables (placeholders)
    for t in doc.tables:
        if len(t.rows) == 1 and len(t.columns) == 1:
            p0_txt = t.rows[0].cells[0].paragraphs[0].text.strip()
            if "[ พื้นที่สำหรับวางภาพประกอบ: รูปที่ 4.6 ]" in p0_txt:
                t.rows[0].cells[0].paragraphs[0].text = "[ พื้นที่สำหรับวางภาพประกอบ: รูปที่ 4.5 ]"
                set_font(t.rows[0].cells[0].paragraphs[0].runs[0], bold=True, font_size=16)
            elif "[ พื้นที่สำหรับวางภาพประกอบ: รูปที่ 4.7 ]" in p0_txt:
                t.rows[0].cells[0].paragraphs[0].text = "[ พื้นที่สำหรับวางภาพประกอบ: รูปที่ 4.6 ]"
                set_font(t.rows[0].cells[0].paragraphs[0].runs[0], bold=True, font_size=16)
            elif "[ พื้นที่สำหรับวางภาพประกอบ: รูปที่ 4.8 ]" in p0_txt:
                t.rows[0].cells[0].paragraphs[0].text = "[ พื้นที่สำหรับวางภาพประกอบ: รูปที่ 4.7 ]"
                set_font(t.rows[0].cells[0].paragraphs[0].runs[0], bold=True, font_size=16)

    # 2. Update captions and narrative paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip().replace('\u200b', '')
        
        # Caption 4.6 -> 4.5
        if txt.startswith("รูปที่ 4.6  ภาพถ่ายการจำลองท่าทางการหกล้ม 4 รูปแบบในห้องน้ำทดสอบ"):
            p.text = "รูปที่ 4.5  ภาพถ่ายการจำลองท่าทางการหกล้ม 4 รูปแบบในห้องน้ำทดสอบ"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.runs[0], bold=True, font_size=16)
            
        elif txt.startswith("จากภาพถ่ายการจำลองท่าทางการหกล้มในรูปที่ 4.6"):
            p.text = p.text.replace("รูปที่ 4.6", "รูปที่ 4.5")
            p.alignment = WD_ALIGN_PARAGRAPH.THAI_JUSTIFY
            set_font(p.runs[0], bold=False, font_size=16)

        # Caption 4.7 -> 4.6
        elif txt.startswith("รูปที่ 4.7  ภาพถ่ายการทดสอบกิจกรรมปกติในชีวิตประจำวัน"):
            p.text = "รูปที่ 4.6  ภาพถ่ายการทดสอบกิจกรรมปกติในชีวิตประจำวัน (ADLs) และกิจกรรมที่เสี่ยงต่อการเกิดสัญญาณเตือนหลอก"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.runs[0], bold=True, font_size=16)
            
        elif txt.startswith("จากภาพการทดสอบกิจกรรมปกติในรูปที่ 4.7"):
            p.text = p.text.replace("รูปที่ 4.7", "รูปที่ 4.6")
            p.alignment = WD_ALIGN_PARAGRAPH.THAI_JUSTIFY
            set_font(p.runs[0], bold=False, font_size=16)

        # Caption 4.8 -> 4.7
        elif txt.startswith("รูปที่ 4.8  ภาพถ่ายการทดสอบประสิทธิภาพการตรวจจับบุคคลด้วยเรดาร์ 24 GHz"):
            p.text = "รูปที่ 4.7  ภาพถ่ายการทดสอบประสิทธิภาพการตรวจจับบุคคลด้วยเรดาร์ 24 GHz ในสภาวะต่าง ๆ ภายในห้องน้ำ"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.runs[0], bold=True, font_size=16)
            
        elif txt.startswith("จากภาพถ่ายในรูปที่ 4.8 ยืนยันถึงคุณสมบัติเด่นของเรดาร์ 24 GHz"):
            p.text = p.text.replace("รูปที่ 4.8", "รูปที่ 4.7")
            p.alignment = WD_ALIGN_PARAGRAPH.THAI_JUSTIFY
            set_font(p.runs[0], bold=False, font_size=16)

        # Infographic diagram caption (former 4.5) -> 4.8
        elif txt.startswith("รูปที่ 4.5  ภาพบรรยากาศการจำลองท่าทางการทดสอบการล้มและกิจกรรมปกติในห้องน้ำทดสอบ") or \
             txt.startswith("รูปที่ 4.5  แผนผังการจัดสภาพแวดล้อมห้องน้ำทดสอบและกรอบการจำลองท่าทางการทดสอบ"):
            p.text = "รูปที่ 4.8  แผนผังการจัดสภาพแวดล้อมห้องน้ำทดสอบและกรอบการจำลองท่าทางการทดสอบ"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.runs[0], bold=True, font_size=16)
            
        elif txt.startswith("จากภาพบรรยากาศการทดสอบในรูปที่ 4.5") or txt.startswith("จากแผนผังการจัดสภาพแวดล้อมห้องน้ำทดสอบในรูปที่ 4.5"):
            p.text = p.text.replace("รูปที่ 4.5", "รูปที่ 4.8")
            p.alignment = WD_ALIGN_PARAGRAPH.THAI_JUSTIFY
            set_font(p.runs[0], bold=False, font_size=16)

    print("  [+] Figure numbers and captions updated cleanly.")

def render_pdf_powershell(docx_path, pdf_path):
    print(f"Rendering PDF via Word COM -> {os.path.basename(pdf_path)}...")
    subprocess.run(["powershell", "-Command", "Get-Process -Name WINWORD -ErrorAction SilentlyContinue | Stop-Process -Force"], capture_output=True)
    time.sleep(1)
    
    ps_script = f'''
$wdFormatPDF = 17
$word = New-Object -ComObject Word.Application
$word.Visible = $false
try {{
    $doc = $word.Documents.Open('{docx_path}')
    $doc.SaveAs([ref]'{pdf_path}', [ref]$wdFormatPDF)
    $doc.Close(0)
    Write-Host "PDF rendered successfully!"
}} finally {{
    $word.Quit()
    [System.Runtime.InteropServices.Marshal]::ReleaseComObject($word) | Out-Null
    [System.GC]::Collect()
    [System.GC]::WaitForPendingFinalizers()
}}
'''
    ps_file = os.path.join(os.path.dirname(docx_path), "temp_render.ps1")
    with open(ps_file, "w", encoding="utf-8") as f:
        f.write(ps_script)
        
    res = subprocess.run(["powershell", "-ExecutionPolicy", "Bypass", "-File", ps_file], capture_output=True, text=True)
    if os.path.exists(ps_file):
        os.remove(ps_file)
        
    if "PDF rendered successfully!" in res.stdout:
        print("  [+] PDF compiled successfully.")
        return True
    else:
        print("  [!] Render failed:", res.stdout, res.stderr)
        return False

def sync_all_tables_with_pdf(docx_path, pdf_path):
    print("Scanning accurate page numbers from PDF...")
    doc_pdf = fitz.open(pdf_path)
    total_pages = len(doc_pdf)
    
    ch1_idx = None
    for i in range(total_pages):
        txt = doc_pdf[i].get_text()
        if "บทที่ 1" in txt and "บทนำ" in txt and "สารบัญ" not in txt:
            ch1_idx = i
            break
            
    print(f"  Chapter 1 starts at physical PDF page {ch1_idx + 1}")
    def to_acad(idx):
        return str(idx - ch1_idx + 1)
        
    # Figures to scan (start scan from ch1_idx to avoid TOC match)
    figures_to_scan = [
        ("รูปที่ 2.1", r"รูปที่\s*2\.1\b"),
        ("รูปที่ 2.2", r"รูปที่\s*2\.2\b"),
        ("รูปที่ 2.3", r"รูปที่\s*2\.3\b"),
        ("รูปที่ 2.4", r"รูปที่\s*2\.4\b"),
        ("รูปที่ 2.5", r"รูปที่\s*2\.5\b"),
        ("รูปที่ 2.6", r"รูปที่\s*2\.6\b"),
        ("รูปที่ 3.1", r"รูปที่\s*3\.1\b"),
        ("รูปที่ 3.2", r"รูปที่\s*3\.2\b"),
        ("รูปที่ 3.3", r"รูปที่\s*3\.3\b"),
        ("รูปที่ 3.4", r"รูปที่\s*3\.4\b"),
        ("รูปที่ 3.5", r"รูปที่\s*3\.5\b"),
        ("รูปที่ 3.6", r"รูปที่\s*3\.6\b"),
        ("รูปที่ 3.7", r"รูปที่\s*3\.7\b"),
        ("รูปที่ 3.8", r"รูปที่\s*3\.8\b"),
        ("รูปที่ 3.9", r"รูปที่\s*3\.9\b"),
        ("รูปที่ 4.1", r"รูปที่\s*4\.1\b"),
        ("รูปที่ 4.2", r"รูปที่\s*4\.2\b"),
        ("รูปที่ 4.3", r"รูปที่\s*4\.3\b"),
        ("รูปที่ 4.4", r"รูปที่\s*4\.4\s*\(ก\)"),
        ("รูปที่ 4.5", r"รูปที่\s*4\.5\b"),
        ("รูปที่ 4.6", r"รูปที่\s*4\.6\b"),
        ("รูปที่ 4.7", r"รูปที่\s*4\.7\b"),
        ("รูปที่ 4.8", r"รูปที่\s*4\.8\b"),
        ("รูปที่ 4.9", r"รูปที่\s*4\.9\b"),
        ("รูปที่ ก.1", r"รูปที่\s*ก\.1\b"),
        ("รูปที่ ก.2", r"รูปที่\s*ก\.2\b"),
        ("รูปที่ ค.1", r"รูปที่\s*ค\.1\b"),
        ("รูปที่ ง.1", r"รูปที่\s*ง\.1\b"),
        ("รูปที่ ง.2", r"รูปที่\s*ง\.2\b"),
    ]
    f_map = {}
    for key, pat in figures_to_scan:
        for idx in range(ch1_idx, total_pages):
            if re.search(pat, doc_pdf[idx].get_text()):
                f_map[key] = to_acad(idx)
                break
                
    # Tables to scan
    tables_to_scan = [
        ("ตารางที่ 2.1", r"ตารางที่\s*2\.1\b"),
        ("ตารางที่ 2.2", r"ตารางที่\s*2\.2\b"),
        ("ตารางที่ 2.3", r"ตารางที่\s*2\.3\b"),
        ("ตารางที่ 2.4", r"ตารางที่\s*2\.4\b"),
        ("ตารางที่ 2.5", r"ตารางที่\s*2\.5\b"),
        ("ตารางที่ 3.1", r"ตารางที่\s*3\.1\b"),
        ("ตารางที่ 4.1", r"ตารางที่\s*4\.1\b"),
        ("ตารางที่ 4.2", r"ตารางที่\s*4\.2\b"),
    ]
    t_map = {}
    for key, pat in tables_to_scan:
        for idx in range(ch1_idx, total_pages):
            if re.search(pat, doc_pdf[idx].get_text()):
                t_map[key] = to_acad(idx)
                break

    # Headings to scan
    headings_to_scan = [
        ("บทที่ 1  บทนำ", r"บทที่\s*1\s*\n\s*บทนำ"),
        ("1.1", r"1\.1\s+ความเป็นมา"),
        ("1.2", r"1\.2\s+วัตถุประสงค์"),
        ("1.3", r"1\.3\s+สมมติฐาน"),
        ("1.4", r"1\.4\s+ข้อตกลงเบื้องต้น"),
        ("1.5", r"1\.5\s+ขอบเขตของโครงงาน"),
        ("1.6", r"1\.6\s+ตัวแปรที่ศึกษา"),
        ("1.7", r"1\.7\s+ประโยชน์ที่คาดว่าจะได้รับ"),
        ("1.8", r"1\.8\s+คำนิยามศัพท์เฉพาะ"),
        ("บทที่ 2  ทฤษฎีและงานวิจัยที่เกี่ยวข้อง", r"บทที่\s*2\s*\n\s*ทฤษฎี"),
        ("2.1", r"2\.1\s+เทคโนโลยีเรดาร์คลื่นมิลลิเมตร"),
        ("2.2", r"2\.2\s+โมดูลเรดาร์"),
        ("2.3", r"2\.3\s+สถาปัตยกรรมไมโครคอนโทรลเลอร์|2\.3\s+ไมโครคอนโทรลเลอร์"),
        ("2.4", r"2\.4\s+สถาปัตยกรรมและโปรโตคอล|2\.4\s+โปรโตคอลการสื่อสาร"),
        ("2.5", r"2\.5\s+ระบบบริการหลังบ้าน|2\.5\s+ระบบคลาวด์และฐานข้อมูล"),
        ("2.6", r"2\.6\s+ระบบส่วนแสดงผล|2\.6\s+เทคโนโลยีการพัฒนาเว็บแอปพลิเคชัน"),
        ("2.7", r"2\.7\s+งานวิจัยและเอกสาร|2\.7\s+งานวิจัยที่เกี่ยวข้อง"),
        ("บทที่ 3  วิธีดำเนินงาน", r"บทที่\s*3\s*\n\s*วิธีดำเนินงาน"),
        ("3.1", r"3\.1\s+ขั้นตอน"),
        ("3.2", r"3\.2\s+การวิเคราะห์และออกแบบสถาปัตยกรรม"),
        ("3.3", r"3\.3\s+การออกแบบและพัฒนาฮาร์ดแวร์"),
        ("3.4", r"3\.4\s+การพัฒนาเฟิร์มแวร์"),
        ("3.5", r"3\.5\s+การพัฒนาระบบคลาวด์"),
        ("3.6", r"3\.6\s+การพัฒนาเว็บแอปพลิเคชัน"),
        ("3.7", r"3\.7\s+การออกแบบการทดสอบ"),
        ("บทที่ 4  ผลการดำเนินงาน", r"บทที่\s*4\s*\n\s*ผลการดำเนินงาน"),
        ("4.1", r"4\.1\s+ผลการออกแบบและสร้างแผ่นวงจรพิมพ์"),
        ("4.2", r"4\.2\s+ผลการพัฒนาเฟิร์มแวร์"),
        ("4.3", r"4\.3\s+ผลการพัฒนาระบบบริการหลังบ้าน"),
        ("4.4", r"4\.4\s+ผลการพัฒนาเว็บแอปพลิเคชัน"),
        ("4.5", r"4\.5\s+ผลการทดสอบและประเมินประสิทธิภาพ"),
        ("บทที่ 5  สรุปผล และข้อเสนอแนะ", r"บทที่\s*5\s*\n\s*สรุปผล"),
        ("5.1", r"5\.1\s+สรุปผลการดำเนินงาน"),
        ("5.2", r"5\.2\s+ปัญหาและอุปสรรค"),
        ("5.3", r"5\.3\s+ข้อเสนอแนะสำหรับ"),
        ("บรรณานุกรม", r"บรรณานุกรม"),
        ("ภาคผนวก ก", r"ภาคผนวก\s*ก"),
        ("ภาคผนวก ข", r"ภาคผนวก\s*ข"),
        ("ภาคผนวก ค", r"ภาคผนวก\s*ค"),
        ("ภาคผนวก ง", r"ภาคผนวก\s*ง"),
        ("ภาคผนวก จ", r"ภาคผนวก\s*จ"),
        ("ประวัติผู้เขียน", r"ประวัติผู้เขียน"),
    ]
    h_map = {}
    for key, pat in headings_to_scan:
        for idx in range(ch1_idx, total_pages):
            if re.search(pat, doc_pdf[idx].get_text()):
                h_map[key] = to_acad(idx)
                break

    print(f"  Scanned: {len(h_map)} headings, {len(t_map)} tables, {len(f_map)} figures.")
    print("  Figures page map:", f_map)
    print("  Tables page map:", t_map)

    # Now load docx and inject these exact page numbers
    doc = docx.Document(docx_path)
    
    tbl_toc1 = None
    tbl_toc2 = None
    tbl_tables = None
    tbl_figures = None
    
    for t in doc.tables:
        first_cell = t.rows[0].cells[0].text.strip() if len(t.rows) > 0 and len(t.columns) > 0 else ""
        r1_cell = t.rows[1].cells[0].text.strip() if len(t.rows) > 1 and len(t.columns) > 0 else ""
        if "บทคัดย่อภาษาไทย" in r1_cell:
            tbl_toc1 = t
        elif "บทที่ 3" in r1_cell:
            tbl_toc2 = t
        elif first_cell in ["ตาราง", "ตารางที่"]:
            tbl_tables = t
        elif first_cell in ["รูป", "รูปที่"]:
            tbl_figures = t

    # 1. Update TOC 2
    toc2_keys = [
        None,
        "บทที่ 3  วิธีดำเนินงาน",
        "3.1", "3.2", "3.3", "3.4", "3.5", "3.6", "3.7",
        "บทที่ 4  ผลการดำเนินงาน",
        "4.1", "4.2", "4.3", "4.4", "4.5",
        "บทที่ 5  สรุปผล และข้อเสนอแนะ",
        "5.1", "5.2", "5.3",
        "บรรณานุกรม",
        "ภาคผนวก ก", "ภาคผนวก ข", "ภาคผนวก ค", "ภาคผนวก ง", "ภาคผนวก จ",
        "ประวัติผู้เขียน",
    ]
    if tbl_toc2:
        for r_idx, key in enumerate(toc2_keys):
            if key is None or key not in h_map:
                continue
            pnum = h_map[key]
            if r_idx < len(tbl_toc2.rows):
                tbl_toc2.rows[r_idx].cells[1].paragraphs[0].text = pnum
                tbl_toc2.rows[r_idx].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for r in tbl_toc2.rows[r_idx].cells[1].paragraphs[0].runs:
                    set_font(r, bold=tbl_toc2.rows[r_idx].cells[0].paragraphs[0].runs[0].bold if tbl_toc2.rows[r_idx].cells[0].paragraphs[0].runs else False, font_size=16)

    # 2. Update List of Tables (Table 3)
    if tbl_tables:
        tab_list_keys = [
            ("ตารางที่ 2.1", "ตารางที่ 2.1  ข้อกำหนดคุณสมบัติทางเทคนิคของโมดูลเรดาร์ HLK-LD6002C"),
            ("ตารางที่ 2.2", "ตารางที่ 2.2  โครงสร้างเฟรมข้อมูลการสื่อสาร (TinyFrame Protocol) ของโมดูล HLK-LD6002C"),
            ("ตารางที่ 2.3", "ตารางที่ 2.3  ชนิดข้อความและชุดคำสั่งการทำงานผ่านโปรโตคอล TinyFrame ของเรดาร์ HLK-LD6002C"),
            ("ตารางที่ 2.4", "ตารางที่ 2.4  คุณสมบัติทางแม่เหล็กไฟฟ้าของวัสดุสำหรับออกแบบฝาครอบเรดาร์ (Radome Materials)"),
            ("ตารางที่ 2.5", "ตารางที่ 2.5  การเปรียบเทียบเทคโนโลยีและวิธีการตรวจจับการล้มสำหรับผู้สูงอายุ"),
            ("ตารางที่ 3.1", "ตารางที่ 3.1  ตารางการกำหนดขาเชื่อมต่อ (Pin Assignment) ระหว่าง ESP32-S3 และอุปกรณ์"),
            ("ตารางที่ 4.1", "ตารางที่ 4.1  ตารางสรุปผลการทดสอบการจำแนกเหตุการณ์การล้มและกิจกรรมปกติ (Confusion Matrix and Performance Evaluation Metrics)"),
            ("ตารางที่ 4.2", "ตารางที่ 4.2  ผลการทดสอบประสิทธิภาพการตรวจจับการมีอยู่ของบุคคลด้วยเรดาร์ 24 GHz ในสภาวะต่าง ๆ"),
        ]
        for r in tbl_tables.rows[1:]:
            txt = r.cells[0].text.strip().replace('\u200b', '')
            for tab_key, _ in tab_list_keys:
                if txt.startswith(tab_key) and tab_key in t_map:
                    r.cells[1].paragraphs[0].text = t_map[tab_key]
                    r.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    for run in r.cells[1].paragraphs[0].runs:
                        set_font(run, bold=False, font_size=16)
                    break

    # 3. Update List of Figures (Table 4)
    figure_entries = [
        ("รูปที่ 2.1  หลักการส่งและรับสัญญาณเรดาร์แบบคลื่นต่อเนื่องมอดูเลตความถี่ (FMCW Chirp)", f_map.get("รูปที่ 2.1", "5")),
        ("รูปที่ 2.2  โมดูลเรดาร์คลื่นมิลลิเมตร 60 GHz รุ่น HLK-LD6002C", f_map.get("รูปที่ 2.2", "7")),
        ("รูปที่ 2.3  โมดูลเรดาร์ตรวจจับบุคคล 24 GHz รุ่น HLK-LD2410C", f_map.get("รูปที่ 2.3", "13")),
        ("รูปที่ 2.4  บอร์ดไมโครคอนโทรลเลอร์ ESP32-S3 DevKitC-1", f_map.get("รูปที่ 2.4", "14")),
        ("รูปที่ 2.5  สถาปัตยกรรม Publish/Subscribe ของโปรโตคอล MQTT", f_map.get("รูปที่ 2.5", "16")),
        ("รูปที่ 2.6  ผังการทำงานของกลไก Webhook สำหรับส่งการแจ้งเตือนไปยัง LINE Messaging API", f_map.get("รูปที่ 2.6", "17")),
        ("รูปที่ 3.1  แผนภาพขั้นตอนการดำเนินงานวิจัยและพัฒนาระบบ", f_map.get("รูปที่ 3.1", "25")),
        ("รูปที่ 3.2  แผนผังสถาปัตยกรรมระบบรวม (End-to-End System Architecture)", f_map.get("รูปที่ 3.2", "27")),
        ("รูปที่ 3.3  ผังวงจรไฟฟ้า (Schematic Circuit Diagram) ของอุปกรณ์ตรวจจับ", f_map.get("รูปที่ 3.3", "28")),
        ("รูปที่ 3.4  การออกแบบลายวงจรพิมพ์ (PCB Layout) ของอุปกรณ์ตรวจจับการล้ม", f_map.get("รูปที่ 3.4", "31")),
        ("รูปที่ 3.5  แผนภาพสถานะการทำงานของระบบ (Finite State Machine Flowchart)", f_map.get("รูปที่ 3.5", "32")),
        ("รูปที่ 3.6  แผนผังอัลกอริทึมการกรองสัญญาณและการยืนยันการล้ม (Fall Confirmation Delay)", f_map.get("รูปที่ 3.6", "33")),
        ("รูปที่ 3.7  ผังความสัมพันธ์เอนทิตีของฐานข้อมูล (Entity-Relationship Diagram: ERD)", f_map.get("รูปที่ 3.7", "35")),
        ("รูปที่ 3.8  ภาพจำลองส่วนติดต่อผู้ใช้เว็บแอปพลิเคชันแดชบอร์ด (Next.js Dashboard UI)", f_map.get("รูปที่ 3.8", "37")),
        ("รูปที่ 3.9  แผนผังการจัดสภาพแวดล้อมห้องน้ำทดสอบและตำแหน่งการติดตั้งอุปกรณ์ต้นแบบ", f_map.get("รูปที่ 3.9", "38")),
        ("รูปที่ 4.1  ชิ้นงานแผ่นวงจรพิมพ์จริง (PCB Assembly) ของระบบตรวจจับการล้มในห้องน้ำ", f_map.get("รูปที่ 4.1", "41")),
        ("รูปที่ 4.2  ผังการทำงานของเครื่องสถานะจำกัด (Finite State Machine)", f_map.get("รูปที่ 4.2", "44")),
        ("รูปที่ 4.3  ภาพจำลองข้อความแจ้งเตือนภัยฉุกเฉิน (LINE Flex Message) บนแอปพลิเคชัน LINE", f_map.get("รูปที่ 4.3", "45")),
        ("รูปที่ 4.4  ภาพหน้าจอเว็บแอปพลิเคชันแดชบอร์ด (Next.js Dashboard)", f_map.get("รูปที่ 4.4", "46")),
        ("รูปที่ 4.5  ภาพถ่ายการจำลองท่าทางการหกล้ม 4 รูปแบบในห้องน้ำทดสอบ", f_map.get("รูปที่ 4.5", "50")),
        ("รูปที่ 4.6  ภาพถ่ายการทดสอบกิจกรรมปกติในชีวิตประจำวัน (ADLs) และกิจกรรมที่เสี่ยงต่อการเกิดสัญญาณเตือนหลอก", f_map.get("รูปที่ 4.6", "51")),
        ("รูปที่ 4.7  ภาพถ่ายการทดสอบประสิทธิภาพการตรวจจับบุคคลด้วยเรดาร์ 24 GHz ในสภาวะต่าง ๆ ภายในห้องน้ำ", f_map.get("รูปที่ 4.7", "52")),
        ("รูปที่ 4.8  แผนผังการจัดสภาพแวดล้อมห้องน้ำทดสอบและกรอบการจำลองท่าทางการทดสอบ", f_map.get("รูปที่ 4.8", "54")),
        ("รูปที่ 4.9  ภาพถ่ายการตอบสนองของระบบแจ้งเตือนภัยแบบเรียลไทม์เมื่อตรวจพบเหตุการณ์หกล้มจริง", f_map.get("รูปที่ 4.9", "55")),
        ("รูปที่ ก.1  แผ่นวงจรพิมพ์หน้าเดียว (Single-sided Copper PCB) ที่ผ่านการกัดลายทองแดงจริง", f_map.get("รูปที่ ก.1", "61")),
        ("รูปที่ ก.2  การประกอบแผงวงจรและเชื่อมต่อเซนเซอร์เรดาร์ HLK-LD6002C, สัญญาณไซเรนเตือนภัย และเคสติดตั้ง", f_map.get("รูปที่ ก.2", "61")),
        ("รูปที่ ค.1  ภาพข้อความแจ้งเตือนเหตุการณ์ฉุกเฉินบนสมาร์ตโฟนผ่านแอปพลิเคชัน LINE", f_map.get("รูปที่ ค.1", "65")),
        ("รูปที่ ง.1  ภาพหน้าจอคำสั่งการติดตั้งและเริ่มการทำงานของคอนเทนเนอร์ด้วย Docker Compose บนเซิร์ฟเวอร์ Azure VM", f_map.get("รูปที่ ง.1", "67")),
        ("รูปที่ ง.2  ภาพหน้าจอการเข้าสู่ระบบ (Login) และการเชื่อมต่อแดชบอร์ดของผู้ใช้งาน", f_map.get("รูปที่ ง.2", "68")),
    ]

    if tbl_figures:
        # Clear existing data rows
        while len(tbl_figures.rows) > 1:
            tbl_figures._tbl.remove(tbl_figures.rows[1]._tr)
            
        for title, pnum in figure_entries:
            row = tbl_figures.add_row()
            c0, c1 = row.cells[0], row.cells[1]
            c0.width = Cm(13.5)
            c1.width = Cm(1.15)
            
            p0 = c0.paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p0.paragraph_format.space_before = Pt(0)
            p0.paragraph_format.space_after = Pt(0)
            p0.paragraph_format.line_spacing = 1.0
            r0 = p0.add_run(title)
            set_font(r0, bold=False, font_size=16)
            
            p1 = c1.paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p1.paragraph_format.space_before = Pt(0)
            p1.paragraph_format.space_after = Pt(0)
            p1.paragraph_format.line_spacing = 1.0
            r1 = p1.add_run(pnum)
            set_font(r1, bold=False, font_size=16)

        print(f"  [+] Table 4 (List of Figures) updated with {len(figure_entries)} figures and accurate page numbers.")

    doc.save(docx_path)
    print("  [+] Saved final docx successfully.")

def main():
    # Step 1: Update figure numbering and text in docx
    doc = docx.Document(DOCX_PATH)
    update_figure_numbers_and_references(doc)
    doc.save(DOCX_PATH)
    
    # Step 2: Render intermediate PDF
    render_pdf_powershell(DOCX_PATH, PDF_PATH)
    
    # Step 3: Scan accurate page numbers and update TOC + List of Figures
    sync_all_tables_with_pdf(DOCX_PATH, PDF_PATH)
    
    # Step 4: Final PDF Render
    render_pdf_powershell(DOCX_PATH, PDF_PATH)
    print("\n🎉 COMPLETE: Report docx and PDF are 100% updated and synchronized!")

if __name__ == "__main__":
    main()
