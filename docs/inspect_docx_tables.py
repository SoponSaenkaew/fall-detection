import docx

doc = docx.Document(r"D:\Final project\fall-detection\docs\เอกสารประกอบโปรเจกต์จบ\รายงานเชิงเทคนิก\รายงานเชิงเทคนิค.docx")

print(f"Total tables: {len(doc.tables)}")

for idx, tbl in enumerate(doc.tables):
    first_cell = tbl.rows[0].cells[0].text.strip() if tbl.rows and tbl.rows[0].cells else ""
    print(f"\n--- TABLE {idx}: first_cell='{first_cell[:40]}' ({len(tbl.rows)} rows, {len(tbl.columns)} cols) ---")
    if "ตารางที่ 4" in first_cell or "สภาวะ" in first_cell or "ประเภทการทดสอบ" in first_cell or idx in [18, 19, 17, 16]:
        for r_i, row in enumerate(tbl.rows):
            row_vals = [c.text.strip().replace('\n', ' ') for c in row.cells]
            print(f"  Row {r_i}: {' | '.join(row_vals)}")
