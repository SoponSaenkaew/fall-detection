import docx, re

doc = docx.Document('docs/เอกสารประกอบโปรเจกต์จบ/รายงานเชิงเทคนิก/รายงานเชิงเทคนิค.docx')

print("=== 1. CHECK CHAPTER FIRST PAGES & SPACING ===")
for i, p in enumerate(doc.paragraphs):
    if re.match(r'^บทที่\s+[1-5]', p.text.strip()) and len(p.text.strip()) < 15:
        print(f"P{i}: '{p.text.strip()}' | align={p.alignment} | sb={p.paragraph_format.space_before} | sa={p.paragraph_format.space_after}")
        if i + 1 < len(doc.paragraphs):
            p1 = doc.paragraphs[i+1]
            print(f"  P{i+1}: '{p1.text.strip()}' | align={p1.alignment} | sb={p1.paragraph_format.space_before} | sa={p1.paragraph_format.space_after}")
        if i + 2 < len(doc.paragraphs):
            p2 = doc.paragraphs[i+2]
            print(f"  P{i+2}: '{p2.text.strip()[:40]}' | sb={p2.paragraph_format.space_before}")

print("\n=== 2. CHECK LEVEL 2 HEADINGS (x.x) ===")
for i, p in enumerate(doc.paragraphs):
    if re.match(r'^[1-5]\.[1-9]\s+', p.text.strip()):
        left = p.paragraph_format.left_indent.cm if p.paragraph_format.left_indent else 0
        first = p.paragraph_format.first_line_indent.cm if p.paragraph_format.first_line_indent else 0
        bolds = [r.bold for r in p.runs]
        sb = p.paragraph_format.space_before.pt if p.paragraph_format.space_before else 0
        sa = p.paragraph_format.space_after.pt if p.paragraph_format.space_after else 0
        print(f"P{i}: '{p.text.strip()[:35]}' | left={left:.2f} | first={first:.2f} | bold={set(bolds)} | sb={sb} | sa={sa}")

print("\n=== 3. CHECK LEVEL 3 HEADINGS (x.x.x) ===")
for i, p in enumerate(doc.paragraphs):
    if re.match(r'^[1-5]\.\d+\.\d+\s+', p.text.strip()):
        left = p.paragraph_format.left_indent.cm if p.paragraph_format.left_indent else 0
        first = p.paragraph_format.first_line_indent.cm if p.paragraph_format.first_line_indent else 0
        bolds = [r.bold for r in p.runs]
        sb = p.paragraph_format.space_before.pt if p.paragraph_format.space_before else 0
        sa = p.paragraph_format.space_after.pt if p.paragraph_format.space_after else 0
        print(f"P{i}: '{p.text.strip()[:35]}' | left={left:.2f} | first={first:.2f} | bold={set(bolds)} | sb={sb} | sa={sa}")

print("\n=== 4. CHECK TABLES & CAPTIONS ===")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith('ตารางที่'):
        left = p.paragraph_format.left_indent.cm if p.paragraph_format.left_indent else 0
        bolds = [r.bold for r in p.runs]
        print(f"P{i}: '{p.text.strip()[:40]}' | left={left:.2f} | bold={set(bolds)} | align={p.alignment}")

print("\n=== 5. CHECK FIGURES & CAPTIONS ===")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith('รูปที่'):
        left = p.paragraph_format.left_indent.cm if p.paragraph_format.left_indent else 0
        bolds = [r.bold for r in p.runs]
        print(f"P{i}: '{p.text.strip()[:40]}' | align={p.alignment} | bold={set(bolds)}")
