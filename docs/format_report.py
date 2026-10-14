import os
import shutil
import docx
from copy import deepcopy
from docx.shared import Cm, Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

BACKUP_PATH = r"D:\Final project\fall-detection\docs\เอกสารประกอบโปรเจกต์จบ\รายงานเชิงเทคนิก\รายงานเชิงเทคนิค_backup.docx"
TARGET_PATH = r"D:\Final project\fall-detection\docs\เอกสารประกอบโปรเจกต์จบ\รายงานเชิงเทคนิก\รายงานเชิงเทคนิค.docx"
DOCS_DIR = r"D:\Final project\fall-detection\docs\เอกสารประกอบโปรเจกต์จบ"
LOGO_PATH = os.path.join(DOCS_DIR, r"คู่มือรูปเล่ม\rmuti_logo.png")

PHOTO_PATH = os.path.join(DOCS_DIR, "รายงานเชิงเทคนิก", "author_photo.jpg")

def remove_borders(table):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'<w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'<w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def set_cell_margins(cell, top=0, bottom=0, left=0, right=0):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def format_vitae_section(doc):
    print("   Formatting Vitae Section with Author Photo...")
    # 1. Locate the two 'ประวัติผู้เขียน' headings
    vitae_paragraphs = []
    for idx, p in enumerate(doc.paragraphs):
        txt = p.text.replace('\u200b', '').strip()
        if txt == "ประวัติผู้เขียน":
            vitae_paragraphs.append((idx, p))
            
    print(f"   Found {len(vitae_paragraphs)} 'ประวัติผู้เขียน' headings.")
    if len(vitae_paragraphs) >= 2:
        # P_sep is the separator sheet (รูปที่ 25: จัดกึ่งกลางหน้ากระดาษ)
        p_sep = vitae_paragraphs[0][1]
        p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sep.paragraph_format.space_before = Pt(250) # Vertically centered
        p_sep.paragraph_format.space_after = Pt(0)
        p_sep.paragraph_format.line_spacing = 1.0
        for r in p_sep.runs:
            set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=True)
            
        # P_bio is the actual biography title (รูปที่ 26: 3.81 ซม. จากขอบบน)
        p_bio = vitae_paragraphs[1][1]
        p_bio.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_bio.paragraph_format.space_before = Pt(0) # 3.81 cm already from top margin
        p_bio.paragraph_format.space_after = Pt(24)
        p_bio.paragraph_format.line_spacing = 1.0
        for r in p_bio.runs:
            set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=True)

    # 2. Find old Sopon table
    target_tbl = None
    for tbl in doc.tables:
        txt = ' '.join(c.text.replace('\u200b', '') for row in tbl.rows for c in row.cells)
        if 'โสภณ' in txt and 'แสนแก้ว' in txt:
            target_tbl = tbl
            break

    if target_tbl is None:
        print("   WARNING: Could not find Sopon table in doc!")
        return

    old_elem = target_tbl._tbl
    parent = old_elem.getparent()
    pos = parent.index(old_elem)

    # 3. Create Master Top Table (Personal info on Left + Photo on Right)
    t_top = doc.add_table(rows=1, cols=2)
    remove_borders(t_top)
    c_left = t_top.cell(0, 0)
    c_right = t_top.cell(0, 1)
    c_left.width = Cm(11.2)
    c_right.width = Cm(3.45)
    set_cell_margins(c_left, 0, 0, 0, 60)
    set_cell_margins(c_right, 0, 0, 60, 0)

    # Photo in right cell
    p_pic = c_right.paragraphs[0]
    p_pic.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_pic.paragraph_format.space_before = Pt(0)
    p_pic.paragraph_format.space_after = Pt(0)
    p_pic.paragraph_format.line_spacing = 1.0
    r_pic = p_pic.add_run()
    r_pic.add_picture(PHOTO_PATH, width=Cm(3.0), height=Cm(4.0))

    # Inner table in left cell
    info_items = [
        ('ชื่อ-นามสกุล', 'นาย... ... (Mr. ...)'),
        ('รหัสนักศึกษา', 'xxxxxxxxxxx-x'),
        ('วัน-เดือน-ปีเกิด', 'xx xxxxxxxx xxxx'),
        ('ที่อยู่', 'xxx ถนน... ตำบล... อำเภอ... จังหวัด... xxxxx'),
        ('อีเมล / ช่องทางติดต่อ', 'user@example.com / GitHub: ...'),
    ]
    t_in = c_left.add_table(rows=len(info_items), cols=2)
    remove_borders(t_in)
    c_left._tc.remove(c_left.paragraphs[0]._p) # remove blank paragraph

    for row_i, (k, v) in enumerate(info_items):
        c0 = t_in.cell(row_i, 0)
        c1 = t_in.cell(row_i, 1)
        c0.width = Cm(3.6)
        c1.width = Cm(7.6)
        set_cell_margins(c0, 0, 0, 0, 40)
        set_cell_margins(c1, 0, 0, 40, 0)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(6)
        p0.paragraph_format.line_spacing = 1.0
        r0 = p0.add_run(k)
        set_run_font(r0, font_name="TH SarabunPSK", font_size=16, bold=False)

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(6)
        p1.paragraph_format.line_spacing = 1.0
        r1 = p1.add_run(v)
        set_run_font(r1, font_name="TH SarabunPSK", font_size=16, bold=False)

    # 4. Education Heading Paragraph
    p_ed_title = doc.add_paragraph()
    p_ed_title.paragraph_format.space_before = Pt(14)
    p_ed_title.paragraph_format.space_after = Pt(8)
    p_ed_title.paragraph_format.line_spacing = 1.0
    r_ed = p_ed_title.add_run('ประวัติการศึกษา')
    set_run_font(r_ed, font_name="TH SarabunPSK", font_size=16, bold=False)

    # 5. Education Table (Indented matching manual p.72)
    tbl_edu = doc.add_table(rows=3, cols=3)
    remove_borders(tbl_edu)

    edu_data = [
        ('วศ.บ.', 'มหาวิทยาลัยเทคโนโลยีราชมงคลอีสาน วิทยาเขตขอนแก่น (กำลังศึกษา)\nสาขาวิชาวิศวกรรมคอมพิวเตอร์ คณะวิศวกรรมศาสตร์'),
        ('ม.6', 'โรงเรียนคำแสนวิทยาสรรค์ ปีการศึกษา 2565\nแผนการเรียนวิทยาศาสตร์-คณิตศาสตร์'),
        ('ม.3', 'โรงเรียนคำแสนวิทยาสรรค์ ปีการศึกษา 2562'),
    ]

    for row_i, (degree, details) in enumerate(edu_data):
        c_indent = tbl_edu.cell(row_i, 0)
        c_deg = tbl_edu.cell(row_i, 1)
        c_det = tbl_edu.cell(row_i, 2)
        c_indent.width = Cm(1.8)
        c_deg.width = Cm(1.8)
        c_det.width = Cm(11.05)
        set_cell_margins(c_indent, 0, 0, 0, 0)
        set_cell_margins(c_deg, 0, 0, 0, 40)
        set_cell_margins(c_det, 0, 0, 40, 0)

        p_deg = c_deg.paragraphs[0]
        p_deg.paragraph_format.space_before = Pt(0)
        p_deg.paragraph_format.space_after = Pt(8)
        p_deg.paragraph_format.line_spacing = 1.0
        r_deg = p_deg.add_run(degree)
        set_run_font(r_deg, font_name="TH SarabunPSK", font_size=16, bold=False)

        p_det = c_det.paragraphs[0]
        p_det.paragraph_format.space_before = Pt(0)
        p_det.paragraph_format.space_after = Pt(8)
        p_det.paragraph_format.line_spacing = 1.0
        r_det = p_det.add_run(details)
        set_run_font(r_det, font_name="TH SarabunPSK", font_size=16, bold=False)

    # Move into position and replace old table
    parent.insert(pos, t_top._tbl)
    parent.insert(pos + 1, p_ed_title._p)
    parent.insert(pos + 2, tbl_edu._tbl)
    parent.remove(old_elem)
    print("   Successfully updated Vitae elements with photo and clean layout.")


def set_run_font(run, font_name="TH SarabunPSK", font_size=16, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)

def set_pgnum(sectPr, fmt=None, start=None):
    pgNum = sectPr.find(qn('w:pgNumType'))
    if pgNum is None:
        pgNum = OxmlElement('w:pgNumType')
        sectPr.append(pgNum)
    if fmt is not None:
        pgNum.set(qn('w:fmt'), fmt)
    elif qn('w:fmt') in pgNum.attrib:
        del pgNum.attrib[qn('w:fmt')]
        
    if start is not None:
        pgNum.set(qn('w:start'), str(start))
    elif qn('w:start') in pgNum.attrib:
        del pgNum.attrib[qn('w:start')]

def fix_document():
    print(f"Restoring clean template from {BACKUP_PATH}...")
    shutil.copyfile(BACKUP_PATH, TARGET_PATH)
    doc = docx.Document(TARGET_PATH)
    
    # =============================================================
    # 1. FORMAT 4 COVER PAGES (ปกนอก, ปกในไทย, ปกในอังกฤษ, ใบรับรอง)
    # =============================================================
    print("\n1. Formatting Cover Pages (ปกนอก, ปกในภาษาไทย, ปกในภาษาอังกฤษ, ใบรับรอง)...")
    
    # --- 1.1 ปกนอก (Outer Cover: P0 - P18) ---
    print("   1.1 Formatting ปกนอก (Outer Cover)...")
    # P0: Logo
    p_logo = doc.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(0)
    p_logo.paragraph_format.space_after = Pt(32) # เว้น 2 บรรทัด (16pt) ตามคู่มือหน้า 42
    
    # P1: Thai title (TH SarabunPSK 20pt ตัวหนา)
    p_title_th = doc.paragraphs[1]
    p_title_th.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title_th.paragraph_format.space_before = Pt(0)
    p_title_th.paragraph_format.space_after = Pt(4)
    p_title_th.paragraph_format.line_spacing = 1.0
    for r in p_title_th.runs:
        set_run_font(r, font_name="TH SarabunPSK", font_size=20, bold=True)
        
    # P2: English title ALL CAPS (TH SarabunPSK 20pt ตัวหนา)
    p_title_en = doc.paragraphs[2]
    p_title_en.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title_en.paragraph_format.space_before = Pt(0)
    p_title_en.paragraph_format.space_after = Pt(0)
    p_title_en.paragraph_format.line_spacing = 1.0
    p_title_en.text = "BATHROOM FALL DETECTION SYSTEM USING 60 GHz MILLIMETER-WAVE RADAR"
    for r in p_title_en.runs:
        set_run_font(r, font_name="TH SarabunPSK", font_size=20, bold=True)
        
    # Remove 5 empty paragraphs between title and author (P3..P7)
    for elem in [doc.paragraphs[i]._element for i in [7, 6, 5, 4, 3]]:
        elem.getparent().remove(elem)

    # Author is now at index 3 (TH SarabunPSK 18pt ตัวหนา, ระยะห่างก่อน-หลังเท่าๆ กัน)
    p_author = doc.paragraphs[3]
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_author.paragraph_format.space_before = Pt(150) # balanced vertical spacing
    p_author.paragraph_format.space_after = Pt(150)
    p_author.paragraph_format.line_spacing = 1.0
    for r in p_author.runs:
        set_run_font(r, font_name="TH SarabunPSK", font_size=18, bold=True)

    # Remove 5 empty paragraphs between author and bottom block (P4..P8)
    for elem in [doc.paragraphs[i]._element for i in [8, 7, 6, 5, 4]]:
        elem.getparent().remove(elem)

    # Bottom block is now at index 4..8 (All 5 lines TH SarabunPSK 16pt ตัวหนา)
    for i in range(4, 9):
        p = doc.paragraphs[i]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2) if i < 8 else Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p_txt = p.text.strip()
        p.text = p_txt
        for r in p.runs:
            set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=True)
            
    # Insert Section Break (Next Page) after Outer Cover (index 8)
    p_b5 = doc.paragraphs[8]
    pPr_b5 = p_b5._element.get_or_add_pPr()
    new_s0_sectPr = deepcopy(doc.sections[0]._sectPr)
    pgMar0 = new_s0_sectPr.find(qn('w:pgMar'))
    if pgMar0 is not None:
        pgMar0.set(qn('w:top'), str(int(2.54 * 567))) # 1440 twips = 2.54 cm
        pgMar0.set(qn('w:bottom'), str(int(2.54 * 567)))
        pgMar0.set(qn('w:left'), str(int(2.54 * 567)))
        pgMar0.set(qn('w:right'), str(int(2.54 * 567)))
    type_elem = new_s0_sectPr.find(qn('w:type'))
    if type_elem is None:
        type_elem = OxmlElement('w:type')
        new_s0_sectPr.append(type_elem)
    type_elem.set(qn('w:val'), 'nextPage')
    pPr_b5.append(new_s0_sectPr)
    print("   Inserted Section Break (Next Page) after Outer Cover.")

    # --- 1.2 ปกในภาษาไทย (Inside Cover Thai: index 9 - 15) ---
    print("   1.2 Formatting ปกในภาษาไทย (Inside Cover Thai)...")
    # Rule: ทั้งหมดใช้ TH SarabunPSK 16pt ตัวบาง (Regular, NOT Bold!)
    # Thai Title only
    p_th_title = doc.paragraphs[9]
    p_th_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_th_title.paragraph_format.space_before = Pt(0)
    p_th_title.paragraph_format.space_after = Pt(0)
    p_th_title.paragraph_format.line_spacing = 1.0
    p_th_title.text = p_th_title.text.strip()
    for r in p_th_title.runs:
        set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=False)

    p_th_author = doc.paragraphs[10]
    p_th_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_th_author.paragraph_format.space_before = Pt(235) # balanced vertical spacing
    p_th_author.paragraph_format.space_after = Pt(235)
    p_th_author.paragraph_format.line_spacing = 1.0
    p_th_author.text = p_th_author.text.strip()
    for r in p_th_author.runs:
        set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=False)

    # Remove empty P11 between author and bottom block
    doc.paragraphs[11]._element.getparent().remove(doc.paragraphs[11]._element)

    # Thai bottom block (index 11..15, all 16pt Regular)
    for i in range(11, 16):
        p = doc.paragraphs[i]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2) if i < 15 else Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p_txt = p.text.strip()
        p.text = p_txt
        for r in p.runs:
            set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=False)

    # Page Break after Thai Inside Cover (index 15)
    run_pb_th = doc.paragraphs[15].add_run()
    run_pb_th.add_break(WD_BREAK.PAGE)
    print("   Added Page Break after Thai Inside Cover.")

    # --- 1.3 ปกในภาษาอังกฤษ (Inside Cover English: index 16 - 23) ---
    print("   1.3 Formatting ปกในภาษาอังกฤษ (Inside Cover English)...")
    # Rule: ทั้งหมดใช้ TH SarabunPSK 16pt ตัวบาง (Regular, NOT Bold!)
    p_en_title = doc.paragraphs[16]
    p_en_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_en_title.paragraph_format.space_before = Pt(0)
    p_en_title.paragraph_format.space_after = Pt(0)
    p_en_title.paragraph_format.line_spacing = 1.0
    p_en_title.text = "Bathroom Fall Detection System Using 60 GHz Millimeter-Wave Radar"
    for r in p_en_title.runs:
        set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=False)

    p_en_author = doc.paragraphs[17]
    p_en_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_en_author.paragraph_format.space_before = Pt(220) # balanced vertical spacing
    p_en_author.paragraph_format.space_after = Pt(220)
    p_en_author.paragraph_format.line_spacing = 1.0
    p_en_author.text = p_en_author.text.strip()
    for r in p_en_author.runs:
        set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=False)

    # English bottom block (index 18..23, all 16pt Regular)
    for i in range(18, 24):
        p = doc.paragraphs[i]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2) if i < 23 else Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p_txt = p.text.strip()
        p.text = p_txt
        for r in p.runs:
            set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=False)

    # Remove empty P24 (was empty P35 in original) and add Page Break to P23
    doc.paragraphs[24]._element.getparent().remove(doc.paragraphs[24]._element)
    run_pb_en = doc.paragraphs[23].add_run()
    run_pb_en.add_break(WD_BREAK.PAGE)
    print("   Added Page Break after English Inside Cover.")

    # --- 1.4 หน้าใบรับรองรายงานเชิงเทคนิค (Approval Certificate) ---
    print("   1.4 Formatting หน้าใบรับรองรายงานเชิงเทคนิค (Approval Certificate)...")
    p_cert_title_elem = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("หัวข้อโครงงาน") and ":" in p.text:
            p_cert_title_elem = p._element
            break

    if p_cert_title_elem is not None:
        # 1. Logo (RMUTI Emblem)
        p_logo2 = doc.add_paragraph()
        p_logo2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo2.paragraph_format.space_before = Pt(0)
        p_logo2.paragraph_format.space_after = Pt(4)
        run_logo2 = p_logo2.add_run()
        run_logo2.add_picture(LOGO_PATH, width=Cm(1.50)) # height ~2.75 cm
        
        # 2. Heading 25pt Bold
        p_h1 = doc.add_paragraph()
        p_h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_h1.paragraph_format.space_before = Pt(0)
        p_h1.paragraph_format.space_after = Pt(2)
        p_h1.paragraph_format.line_spacing = 1.0
        r_h1 = p_h1.add_run("ใบรับรองรายงานเชิงเทคนิค")
        set_run_font(r_h1, font_name="TH SarabunPSK", font_size=25, bold=True)
        
        # 3. University Subtitle 20pt Bold
        p_h2 = doc.add_paragraph()
        p_h2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_h2.paragraph_format.space_before = Pt(0)
        p_h2.paragraph_format.space_after = Pt(14)
        p_h2.paragraph_format.line_spacing = 1.0
        r_h2 = p_h2.add_run("มหาวิทยาลัยเทคโนโลยีราชมงคลอีสาน วิทยาเขตขอนแก่น")
        set_run_font(r_h2, font_name="TH SarabunPSK", font_size=20, bold=True)
        
        p_cert_title_elem.addprevious(p_logo2._element)
        p_cert_title_elem.addprevious(p_h1._element)
        p_cert_title_elem.addprevious(p_h2._element)
        print("   Inserted RMUTI logo and title block into Certificate page.")

    # Spacing and fonts on certificate body
    found_cert = False
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("หัวข้อโครงงาน") and ":" in txt:
            found_cert = True
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.0
            for r in p.runs:
                set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=False)
        elif found_cert:
            if txt.startswith(("จัดทำโดย", "สาขาวิชา", "ประธานที่ปรึกษา")):
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=False)
            elif txt.startswith("ที่ปรึกษาร่วม"):
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(10)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=False)
            elif "ได้รับอนุมัติให้เป็นส่วนหนึ่ง" in txt:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(12)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=False)
            elif "คณบดีคณะวิศวกรรมศาสตร์" in txt or "วันที่........เดือน" in txt:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(3 if "คณบดี" in txt else 12)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=False)
            elif "คณะกรรมการสอบรายงานเชิงเทคนิค" in txt:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=True)
                break

    if doc.tables:
        t0 = doc.tables[0]
        for row in t0.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    for r in p.runs:
                        set_run_font(r, font_name="TH SarabunPSK", font_size=15, bold=r.bold)

    # =============================================================
    # 2. REPLACE 1x1 PLACEHOLDER TABLES WITH IMAGES
    # =============================================================
    print("\n2. Inserting available high-resolution images for figures...")
    img_map = {
        "รูปที่ 2.1": os.path.join(DOCS_DIR, r"รูปภาพทฤษฎี\01_fmcw_radar_principle_detailed.png"),
        "รูปที่ 3.5": os.path.join(DOCS_DIR, r"แฟ้มภาพประกอป\fsm_diagram_academic.png"),
        "รูปที่ 3.6": os.path.join(DOCS_DIR, r"รูปภาพทฤษฎี\12_fall_detection_state_machine.png"),
    }
    
    for t in list(doc.tables):
        if len(t.rows) == 1 and len(t.columns) == 1:
            cell_txt = t.rows[0].cells[0].text.strip().replace('\u200b', '')
            matched_fig = None
            for fig_key in img_map:
                if fig_key in cell_txt:
                    matched_fig = fig_key
                    break
                    
            if matched_fig and os.path.exists(img_map[matched_fig]):
                img_path = img_map[matched_fig]
                print(f"   Replacing placeholder table for {matched_fig} with {os.path.basename(img_path)}...")
                img_p = doc.add_paragraph()
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p.paragraph_format.space_before = Pt(16)
                img_p.paragraph_format.space_after = Pt(0)
                run_img = img_p.add_run()
                run_img.add_picture(img_path, width=Inches(5.6))
                
                t._tbl.addprevious(img_p._p)
                t._tbl.getparent().remove(t._tbl)
                print(f"   Replaced {matched_fig} placeholder with image.")

    # Insert Figure 4.2 flowchart if not already present
    flowchart_path = os.path.join(DOCS_DIR, r"แฟ้มภาพประกอป\firmware_flowchart.png")
    if os.path.exists(flowchart_path):
        for i, p in enumerate(doc.paragraphs):
            text = p.text.strip().replace('\u200b', '')
            if text.startswith("รูปที่ 4.2"):
                prev_has_drawing = 'w:drawing' in doc.paragraphs[i-1]._element.xml if i > 0 else False
                if not prev_has_drawing:
                    print("   Inserting firmware flowchart image for รูปที่ 4.2...")
                    img_p = doc.add_paragraph()
                    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    img_p.paragraph_format.space_before = Pt(16)
                    img_p.paragraph_format.space_after = Pt(0)
                    run_img = img_p.add_run()
                    run_img.add_picture(flowchart_path, width=Inches(5.5))
                    p._p.addprevious(img_p._p)
                    print("   Inserted firmware flowchart image for รูปที่ 4.2.")
                break

    # =============================================================
    # 3. FORMAT TABLE OF CONTENTS (TABLE 1) & INSERT SECTION BREAK BEFORE TOC
    # =============================================================
    print("\n3. Formatting Table of Contents (สารบัญ)...")
    t1 = doc.tables[1]
    row0_txt = [c.text.strip().replace('\u200b', '') for c in t1.rows[0].cells]
    if row0_txt[0] == "สารบัญ":
        t1._tbl.remove(t1.rows[0]._tr)
        print("   Removed redundant title row from Table 1.")

    p_toc_break = doc.add_paragraph()
    pPr_toc = p_toc_break._p.get_or_add_pPr()
    new_toc_sect = deepcopy(doc.sections[-1]._sectPr)
    set_pgnum(new_toc_sect, fmt='thaiLetters', start=None)
    if new_toc_sect.find(qn('w:titlePg')) is None:
        new_toc_sect.append(OxmlElement('w:titlePg'))
    pPr_toc.append(new_toc_sect)
    t1._tbl.addprevious(p_toc_break._p)
    print("   Inserted Section Break before Table of Contents (TOC).")

    p_toc_head = doc.add_paragraph()
    p_toc_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_toc_head.paragraph_format.space_before = Pt(36) # 3.81 + 1.27 = 5.08 cm
    p_toc_head.paragraph_format.space_after = Pt(16)
    run_head = p_toc_head.add_run("สารบัญ")
    set_run_font(run_head, font_name="TH SarabunPSK", font_size=16, bold=True)
    t1._tbl.addprevious(p_toc_head._p)
    print("   Added 'สารบัญ' heading with 5.08 cm top spacing before Table 1.")

    # =============================================================
    # 4. INSERT SECTION BREAK BEFORE CHAPTER 1
    # =============================================================
    print("\n4. Inserting Section Break before Chapter 1...")
    body = doc._body._element
    c1_idx = None
    for i, elem in enumerate(body):
        if elem.tag.endswith('p'):
            p = docx.text.paragraph.Paragraph(elem, doc)
            if p.text.strip().startswith("บทที่ 1"):
                c1_idx = i
                break
                
    if c1_idx is not None:
        elem_before_c1 = body[c1_idx - 1]
        pPr = elem_before_c1.get_or_add_pPr()
        new_sect = deepcopy(doc.sections[-1]._sectPr)
        set_pgnum(new_sect, fmt='thaiLetters', start=None)
        if new_sect.find(qn('w:titlePg')) is None:
            new_sect.append(OxmlElement('w:titlePg'))
        pPr.append(new_sect)
        print("   Inserted Section Break before Chapter 1.")

    # =============================================================
    # 5. INSERT SECTION BREAK BEFORE CHAPTER 3
    # =============================================================
    print("\n5. Inserting Section Break before Chapter 3...")
    c3_elem_idx = None
    for i, elem in enumerate(body):
        if elem.tag.endswith('p'):
            p = docx.text.paragraph.Paragraph(elem, doc)
            if p.text.strip().startswith("บทที่ 3"):
                c3_elem_idx = i
                break

    if c3_elem_idx is not None:
        p_c3_break = doc.add_paragraph()
        pPr3 = p_c3_break._p.get_or_add_pPr()
        new_c3_sect = deepcopy(doc.sections[-1]._sectPr)
        set_pgnum(new_c3_sect, fmt='decimal', start=None)
        if new_c3_sect.find(qn('w:titlePg')) is None:
            new_c3_sect.append(OxmlElement('w:titlePg'))
        pPr3.append(new_c3_sect)
        body[c3_elem_idx].addprevious(p_c3_break._p)
        print("   Inserted clean Section Break before Chapter 3.")

    # Save and reload to let python-docx rebuild doc.sections cleanly
    doc.save(TARGET_PATH)
    doc = docx.Document(TARGET_PATH)

    # =============================================================
    # 6. CONFIGURE PAGE SIZES, MARGINS & PAGINATION PER SECTION
    # =============================================================
    print(f"\n6. Configuring all {len(doc.sections)} sections (A4, Margins, Header/Footer)...")
    for idx, sec in enumerate(doc.sections):
        # A4 standard
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        
        if idx == 0:
            # Section 0: ปกนอก (Outer Cover) -> 2.54 cm all sides
            sec.top_margin = Cm(2.54)
            sec.bottom_margin = Cm(2.54)
            sec.left_margin = Cm(2.54)
            sec.right_margin = Cm(2.54)
            sec.header_distance = Cm(2.54)
            sec.footer_distance = Cm(2.54)
            sec.different_first_page_header_footer = False
        else:
            # Standard thesis margins: Top 3.81, Bottom 2.54, Left 3.81, Right 2.54
            sec.top_margin = Cm(3.81)
            sec.bottom_margin = Cm(2.54)
            sec.left_margin = Cm(3.81)
            sec.right_margin = Cm(2.54)
            sec.header_distance = Cm(2.54)
            sec.footer_distance = Cm(2.54)

    print(f"   Configuring pagination for {len(doc.sections)} sections:")
    # Sec 1: Front covers (ก, ข, ค - hidden)
    set_pgnum(doc.sections[1]._sectPr, fmt='thaiLetters', start=1)
    
    # Sec 2: บทคัดย่อ starts at 4 ('ง' in Thai letters)
    set_pgnum(doc.sections[2]._sectPr, fmt='thaiLetters', start=4)
    doc.sections[2].different_first_page_header_footer = False
    
    # Sec 3: สารบัญ
    set_pgnum(doc.sections[3]._sectPr, fmt='thaiLetters', start=None)
    doc.sections[3].different_first_page_header_footer = True
    
    # Sec 4: สารบัญตาราง
    set_pgnum(doc.sections[4]._sectPr, fmt='thaiLetters', start=None)
    doc.sections[4].different_first_page_header_footer = True
    
    # Sec 5: สารบัญรูป
    set_pgnum(doc.sections[5]._sectPr, fmt='thaiLetters', start=None)
    doc.sections[5].different_first_page_header_footer = True
    
    # Sec 6: คำอธิบายสัญลักษณ์และคำย่อ
    set_pgnum(doc.sections[6]._sectPr, fmt='thaiLetters', start=None)
    doc.sections[6].different_first_page_header_footer = False
    
    # Sec 7: บทที่ 1 -> restarts at decimal 1!
    set_pgnum(doc.sections[7]._sectPr, fmt='decimal', start=1)
    doc.sections[7].different_first_page_header_footer = True
    
    # Sec 8 onwards: decimal continue
    for s_idx in range(8, len(doc.sections)):
        set_pgnum(doc.sections[s_idx]._sectPr, fmt='decimal', start=None)
        doc.sections[s_idx].different_first_page_header_footer = True

    # =============================================================
    # 7. ADJUST INDENTS IN FRONT MATTER
    # =============================================================
    print("\n7. Adjusting front matter indents...")
    for i in range(30, 85):
        if i < len(doc.paragraphs):
            p = doc.paragraphs[i]
            fl = p.paragraph_format.first_line_indent
            if fl and abs(fl.cm - 1.25) < 0.1:
                p.paragraph_format.first_line_indent = Cm(1.80)
                print(f"   P{i} indent updated to 1.8 cm: {p.text[:35]}...")

    # =============================================================
    # 8. STANDARDIZE CHAPTER HEADINGS & 5.08 CM TOP MARGIN (space_before=36pt)
    # =============================================================
    print("\n8. Setting 5.08 cm top margin (space_before=36pt) on chapter headings...")
    major_headings = [
        "สารบัญตาราง", "สารบัญรูป", "คำอธิบายสัญลักษณ์และคำย่อ",
        "บทที่ 1", "บทที่ 2", "บทที่ 3", "บทที่ 4", "บทที่ 5",
        "บรรณานุกรม"
    ]
    
    for p in doc.paragraphs:
        text = p.text.strip().replace('\u200b', '')
        for mh in major_headings:
            if text.startswith(mh):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(36) # 3.81 + 1.27 = 5.08 cm
                p.paragraph_format.space_after = Pt(16)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=True)
                print(f"   Formatted major heading: {text[:35]}")
                break

    # Subtitles for Chapter 3 & 4
    for p in doc.paragraphs:
        text = p.text.strip().replace('\u200b', '')
        if text in ["วิธีดำเนินงาน", "ผลการดำเนินงาน"]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(16)
            p.paragraph_format.line_spacing = 1.0
            for r in p.runs:
                set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=True)

    # =============================================================
    # 9. STANDARDIZE MAIN HEADINGS (x.x) & SUBHEADINGS
    # =============================================================
    print("\n9. Setting 16pt line spacing before main headings (x.x)...")
    for p in doc.paragraphs:
        text = p.text.strip().replace('\u200b', '')
        if len(text) > 3 and text[0].isdigit() and text[1] == '.' and text[2].isdigit() and (len(text) <= 3 or text[3] in [' ', '\t']):
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0)
            for r in p.runs:
                set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=True)

    print("\n10. Ensuring Chapter 5 subheadings (5.2.x, 5.3.x) are Bold...")
    for p in doc.paragraphs:
        text = p.text.strip().replace('\u200b', '')
        if text.startswith(("5.2.1", "5.2.2", "5.2.3", "5.3.1", "5.3.2", "5.3.3")):
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if p.runs:
                set_run_font(p.runs[0], font_name="TH SarabunPSK", font_size=16, bold=True)
            print(f"   Set bold on {text[:25]}")

    # =============================================================
    # 10. STANDARDIZE TABLE & FIGURE CAPTIONS
    # =============================================================
    print("\n11. Formatting table and figure captions...")
    for p in doc.paragraphs:
        text = p.text.strip().replace('\u200b', '')
        if text.startswith("ตารางที่"):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            for r in p.runs:
                set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=True)

        elif text.startswith("รูปที่"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(16)
            p.paragraph_format.line_spacing = 1.0
            for r in p.runs:
                set_run_font(r, font_name="TH SarabunPSK", font_size=16, bold=False)

    # =============================================================
    # 11. GLOBAL PASS ON LINE SPACING (1.0)
    # =============================================================
        # =============================================================
    # 12. FORMAT VITAE SECTION (ประวัติผู้เขียน พร้อมรูปถ่าย)
    # =============================================================
    print("\n12. Formatting Vitae Section with Author Photo...")
    format_vitae_section(doc)

    print("\n13. Global pass on line spacing...")
    for p in doc.paragraphs:
        if p.paragraph_format.line_spacing is not None and p.paragraph_format.line_spacing != 1.0:
            p.paragraph_format.line_spacing = 1.0

    # Save final document
    doc.save(TARGET_PATH)
    print(f"\n==========================================================")
    print(f"SUCCESS: Fully updated and standardized {TARGET_PATH}!")
    print(f"==========================================================")

if __name__ == '__main__':
    fix_document()
