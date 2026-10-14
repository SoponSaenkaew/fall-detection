import docx

doc = docx.Document(r"D:\Final project\fall-detection\docs\เอกสารประกอบโปรเจกต์จบ\รายงานเชิงเทคนิก\รายงานเชิงเทคนิค.docx")

for i in range(345, 375):
    if i < len(doc.paragraphs):
        p = doc.paragraphs[i]
        drawings = p._p.xpath(".//w:drawing")
        print(f"P{i}: drawings={len(drawings)} | {p.text.strip()[:80]}")
